import openpyxl

# 1. Cargar el documento
archivo = "LEGO OPERACIONES AGO 26.xlsx"
wb = openpyxl.load_workbook(archivo)
ws_lista = wb["Lista Elementos"]
ws_info = wb["Elementos Info"]

# Comprobar si existe la hoja de destino, si no, crearla
if "LEGO OPERACIONES" not in wb.sheetnames:
    ws_destino = wb.create_sheet("LEGO OPERACIONES")
else:
    ws_destino = wb["LEGO OPERACIONES"]

# 2. Extraer la lista de nombres requeridos (Cambiando Ñ y ñ por /)
nombres_requeridos = []
for row in range(2, ws_lista.max_row + 1):
    nombre = ws_lista.cell(row=row, column=1).value
    if nombre:
        nombre_modificado = str(nombre).replace("Ñ", "/").replace("ñ", "/")
        nombres_requeridos.append(nombre_modificado.strip().split())

# 3. Mapear los bloques en "Elementos Info" por persona
bloques_info = []
bloque_actual = []
nombre_actual = None

for row in range(1, ws_info.max_row + 1):
    celda_nombre = ws_info.cell(row=row, column=6) # Columna F
    
    # Si la celda tiene valor y la fuente está en negrita, es una nueva persona
    if celda_nombre.value and celda_nombre.font and celda_nombre.font.bold:
        if bloque_actual:
            bloques_info.append({"nombre": nombre_actual, "filas": bloque_actual})
        nombre_actual = str(celda_nombre.value).strip().split()
        bloque_actual = [row]
    elif nombre_actual:
        bloque_actual.append(row)

# Guardar el último bloque leído
if bloque_actual:
    bloques_info.append({"nombre": nombre_actual, "filas": bloque_actual})

# 4. Filtrar y pegar en "LEGO OPERACIONES" a partir de la fila 22
fila_destino = 22
for persona in bloques_info:
    es_requerido = any(all(palabra in persona["nombre"] for palabra in req) for req in nombres_requeridos)
    
    if es_requerido:
        for fila_origen in persona["filas"]:
            for col in range(1, ws_info.max_column + 1):
                valor_origen = ws_info.cell(row=fila_origen, column=col).value
                ws_destino.cell(row=fila_destino, column=col).value = valor_origen
            fila_destino += 1
        fila_destino += 1 # Espacio de separación

# 5. --- COMPROBACIÓN DE ELEMENTOS FALTANTES ---

lista_original = [" ".join(req) for req in nombres_requeridos]
nombres_encontrados = []

for persona in bloques_info:
    for req in nombres_requeridos:
        if all(palabra in persona["nombre"] for palabra in req):
            nombres_encontrados.append(" ".join(req))
            break # Si ya lo encontró, pasa a la siguiente persona

faltantes = set(lista_original) - set(nombres_encontrados)

print(f"\nSe encontraron {len(nombres_encontrados)} elementos.")
if len(faltantes) > 0:
    print(f"Faltan {len(faltantes)} elementos. Los nombres que no hicieron match son:")
    for nombre in faltantes:
        print("-", nombre)
else:
    print("Todo bien. No hubo errores.")

# Guardar
wb.save("LEGO_OPERACIONES_ACTUALIZADO.xlsx")